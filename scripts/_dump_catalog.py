from openpyxl import load_workbook
path = r"E:\Lanza\Projetos\300 bussola\catalogo_chamados_300.xlsx"
wb = load_workbook(path, data_only=True)
for name in wb.sheetnames:
    ws = wb[name]
    print("\n====", name, ws.sheet_state, "====")
    n = 0
    for row in ws.iter_rows(values_only=True):
        if any(c is not None and str(c).strip() for c in row):
            n += 1
            if n <= 40:
                print(" | ".join("" if c is None else str(c).replace("\n"," ")[:70] for c in row[:9]))
    print("NONEMPTY", n)
wb.close()
