import json
from openpyxl import load_workbook

def summarize(path, sample_rows=4, max_cols=18):
    wb = load_workbook(path, read_only=True, data_only=False)
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        state = ws.sheet_state
        rows = []
        for i, row in enumerate(ws.iter_rows(max_row=sample_rows, max_col=max_cols, values_only=True)):
            rows.append([None if c is None else str(c)[:80] for c in row])
        sheets.append({
            "name": name,
            "state": state,
            "dims": ws.dimensions if hasattr(ws, "dimensions") else None,
            "sample": rows,
        })
    wb.close()
    return sheets

adp = r"E:\Lanza\Projetos\300 bussola\ADP BH SCALE_2026.xlsx"
cat = r"E:\Lanza\Projetos\300 bussola\catalogo_chamados_300.xlsx"
out = {
    "adp_sheets": [{"name": s["name"], "state": s["state"]} for s in summarize(adp, sample_rows=0)],
}
# catalog full-ish
wb = load_workbook(cat, data_only=True)
catalog = []
for name in wb.sheetnames:
    ws = wb[name]
    rows = []
    for row in ws.iter_rows(values_only=True):
        if any(c is not None and str(c).strip() for c in row):
            rows.append([None if c is None else str(c) for c in row])
    catalog.append({"name": name, "state": ws.sheet_state, "rows": rows})
wb.close()
out["catalog"] = catalog
path = r"E:\Lanza\Projetos\300 bussola\scripts\_inspect_out.json"
with open(path, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=2)
print("adp", len(out["adp_sheets"]), "catalog", len(catalog))
print("hidden adp", sum(1 for s in out["adp_sheets"] if s["state"] != "visible"))
