from openpyxl import load_workbook
path = r"E:\Lanza\Projetos\300 bussola\ADP BH SCALE_2026.xlsx"
wb = load_workbook(path, data_only=True, read_only=True)
needles = ("1,2", "1.2", "proje", "DATA DA CALL", "SÓCIO", "SOCIO")
hits = {k: [] for k in needles}
for name in wb.sheetnames:
    ws = wb[name]
    for row in ws.iter_rows(max_row=80, max_col=20, values_only=False):
        for cell in row:
            if cell.value is None:
                continue
            s = str(cell.value)
            low = s.lower()
            for n in needles:
                if n.lower() in low or n.lower() in s.lower():
                    if len(hits[n]) < 8:
                        hits[n].append(f"{name}!{cell.coordinate}={s[:80]}")
for k,v in hits.items():
    print("\n##", k, len(v))
    for x in v:
        print(x)
wb.close()
