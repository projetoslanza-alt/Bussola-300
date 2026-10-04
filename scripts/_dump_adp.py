from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

path = r"E:\Lanza\Projetos\300 bussola\ADP BH SCALE_2026.xlsx"
wb = load_workbook(path, data_only=True, read_only=True)

targets = ["ADP GERAL OUT26", "MANÍACOS OUT26", "2809 Á 0210", "RESULTADO ACUMULADO"]
# find similar names
names = wb.sheetnames
print("RESULT matches", [n for n in names if "RESULT" in n.upper() or "ACUM" in n.upper()])

def dump(ws, max_r=45, max_c=16):
    print("\n====", ws.title, "state", ws.sheet_state, "====")
    for r in range(1, max_r + 1):
        cells = []
        for c in range(1, max_c + 1):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip():
                cells.append(f"{get_column_letter(c)}{r}={str(v).replace(chr(10),' ')[:60]}")
        if cells:
            print(f"R{r}: " + " | ".join(cells))

for name in ["ADP GERAL OUT26", "MANÍACOS OUT26", "2809 Á 0210"]:
    dump(wb[name], 35, 14)
wb.close()
